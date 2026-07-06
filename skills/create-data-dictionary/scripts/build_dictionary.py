"""
Build an Excel data dictionary from a Tableau workbook.

Usage:
  python build_dictionary.py WORKBOOK.twbx [-o OUTPUT.xlsx] [-d DESCRIPTIONS.json]

DESCRIPTIONS.json (optional) maps a field caption -> a plain-English description,
e.g. {"Profit Ratio": "Profit as a share of sales."}. Generate these with Claude
after inspecting the resolved formulas (see SKILL.md), then pass them in so the
dictionary reads in business language rather than raw formulas.
"""
import argparse, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from tableau_xml import parse_workbook
from lint_workbook import lint, add_findings_sheet, CHECKS, legend_block
from dependency_graph import PALETTE   # shared field-kind colour palette (no matplotlib import)

def status(f):
    if f["kind"] == "parameter": return "Parameter"
    if f["kind"] == "calculated": return "Unused calc" if f["unused"] else "Active"
    return "Active" if f["used_in"] else "Not placed on a sheet"

def _actions_sheet(book, actions):
    """Add a sheet listing the workbook's dashboard actions (filter / highlight / URL / etc.)."""
    HEAD="2F5496"; arial=lambda **k: Font(name="Arial", **k)
    thin=Side(style="thin",color="D9D9D9"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    sh=book.create_sheet("Actions")
    cols=["Action","Type","Trigger","Source (dashboard · sheet)","Target"]
    TYPE_FILL={"Filter":"DDEBF7","Highlight":"FFF2CC","Go to URL":"E2EFDA","Change parameter":"FCE4D6"}
    for i,w in enumerate([30,16,12,40,34],1):
        sh.column_dimensions[get_column_letter(i)].width=w
    hdr=legend_block(sh, [(TYPE_FILL[t],t) for t in ["Filter","Highlight","Go to URL","Change parameter"]])
    for i in range(1,6):
        c=sh.cell(row=hdr,column=i,value=cols[i-1]); c.font=arial(bold=True,color="FFFFFF",size=10)
        c.fill=PatternFill("solid",fgColor=HEAD); c.alignment=Alignment(vertical="center",wrap_text=True)
    sh.row_dimensions[hdr].height=22
    rn=hdr+1
    for a in actions:
        src=" · ".join(x for x in [a["source_dashboard"],a["source_sheet"]] if x)
        vals=[a["caption"],a["type"],a["trigger"],src,a["target"]]
        fill=TYPE_FILL.get(a["type"],"FFFFFF")
        for ci in range(1,6):
            c=sh.cell(row=rn,column=ci,value=vals[ci-1]); c.font=arial(size=9.5)
            c.alignment=Alignment(vertical="top",wrap_text=True); c.border=border
            if fill!="FFFFFF": c.fill=PatternFill("solid",fgColor=fill)
        rn+=1
    if rn==hdr+1:
        sh.cell(row=rn,column=1,value="— no dashboard actions —"); rn+=1
    sh.freeze_panes=f"A{hdr+1}"; sh.auto_filter.ref=f"A{hdr}:E{max(rn-1,hdr)}"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("-o", "--output")
    ap.add_argument("-d", "--descriptions")
    ap.add_argument("--no-graph", action="store_true",
                    help="skip the Dependency Graph sheet (which is added by default)")
    a = ap.parse_args()
    wb = parse_workbook(a.workbook)
    _, findings = lint(wb)
    desc = json.load(open(a.descriptions)) if a.descriptions else {}
    out = a.output or os.path.splitext(os.path.basename(a.workbook))[0] + "_Data_Dictionary.xlsx"

    fields = sorted(wb["fields"], key=lambda f: ({"parameter":0,"calculated":1,"field":2}[f["kind"]],
                                                 f["datasource"], f["caption"].lower()))
    NAVY="1F3864"; HEAD="2F5496"; WARN=PALETTE["unused"][0]
    KIND_FILL={k:PALETTE[k][0] for k in ("parameter","calculated","field")}
    arial=lambda **k: Font(name="Arial", **k)
    book=Workbook()

    # Summary
    s=book.active; s.title="Summary"
    s["A1"]="Tableau Workbook — Data Dictionary"; s["A1"].font=arial(bold=True,size=15,color="FFFFFF")
    s["A1"].fill=PatternFill("solid",fgColor=NAVY); s.merge_cells("A1:C1"); s.row_dimensions[1].height=24
    s["A2"]=wb["source"]; s["A2"].font=arial(italic=True,size=10,color="555555")
    nk=lambda k: sum(f["kind"]==k for f in fields)
    rows=[("",""),("Datasources",str(len(wb["datasources"]))," · ".join(wb["datasources"])),
          ("Worksheets",str(len(wb["worksheets"])),""),("Dashboards",str(len(wb["dashboards"]))," · ".join(wb["dashboards"])),
          ("Dashboard actions",str(len(wb["actions"]))," · ".join(sorted({a["type"] for a in wb["actions"]}))),
          ("Fields documented",str(len(fields)),""),("  Parameters",str(nk("parameter")),""),
          ("  Calculated fields",str(nk("calculated")),""),("  Physical columns",str(nk("field")),"")]
    r=3
    for row in rows:
        a_,b_,*c_=row; c_=c_[0] if c_ else ""
        s[f"A{r}"]=a_; s[f"B{r}"]=b_; s[f"C{r}"]=c_
        s[f"A{r}"].font=arial(bold=bool(a_ and not a_.startswith(" ")),size=10)
        s[f"B{r}"].font=arial(size=10); s[f"C{r}"].font=arial(size=9,color="555555"); r+=1
    r+=1
    unused=[f for f in fields if f["unused"]]
    s[f"A{r}"]="Calculated fields not used on any sheet and not feeding another field:"
    s[f"A{r}"].font=arial(bold=True,size=10,color="C0504D"); r+=1
    for f in unused:
        s[f"A{r}"]=f"   • {f['caption']}  ({f['datasource']})"; s[f"A{r}"].font=arial(size=10); r+=1
    if not unused:
        s[f"A{r}"]="   • none"; s[f"A{r}"].font=arial(size=10,italic=True); r+=1
    r+=1
    s[f"A{r}"]="Health check (lint):"; s[f"A{r}"].font=arial(bold=True,size=10,color="2F5496"); r+=1
    for label,key,_why,_fill in CHECKS:
        s[f"A{r}"]=f"   {label}"; s[f"B{r}"]=len(findings[key])
        s[f"A{r}"].font=arial(size=10); s[f"B{r}"].font=arial(size=10,bold=True); r+=1
    s[f"A{r}"]="   → see the Lint sheet for details & suggested fixes"
    s[f"A{r}"].font=arial(size=9,italic=True,color="555555")
    s.column_dimensions["A"].width=32; s.column_dimensions["B"].width=10; s.column_dimensions["C"].width=70

    # Dictionary
    d=book.create_sheet("Data Dictionary")
    cols=["Field","Datasource","Kind","Role","Data Type","Description","Formula",
          "Depends On","Feeds Into","Used In Sheets","Status"]
    widths=[26,20,12,11,11,46,52,26,26,34,16]
    for i,w in enumerate(widths,1):
        d.column_dimensions[get_column_letter(i)].width=w
    hdr=legend_block(d, [(KIND_FILL["calculated"],"Calculated field"),(KIND_FILL["parameter"],"Parameter"),
                         (KIND_FILL["field"],"Physical column"),(WARN,"Unused calc (also flagged in Status)")])
    for i in range(1,len(cols)+1):
        c=d.cell(row=hdr,column=i,value=cols[i-1]); c.font=arial(bold=True,color="FFFFFF",size=10)
        c.fill=PatternFill("solid",fgColor=HEAD); c.alignment=Alignment(vertical="center",wrap_text=True)
    d.row_dimensions[hdr].height=28
    thin=Side(style="thin",color="D9D9D9"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    rn=hdr+1
    for f in fields:
        st=status(f)
        dsc=desc.get(f["caption"], "" if f["kind"]!="field" else f"Physical column from {f['datasource']}.")
        vals=[f["caption"],f["datasource"],f["kind"].title(),(f["role"] or "").title(),
              (f["datatype"] or "").title(),dsc,f["formula"],", ".join(f["dependencies"]),
              ", ".join(f["feeds_into"]),", ".join(f["used_in"]),st]
        fill=WARN if st=="Unused calc" else KIND_FILL[f["kind"]]
        for ci in range(1,len(cols)+1):
            c=d.cell(row=rn,column=ci,value=vals[ci-1])
            c.font=Font(name=("Consolas" if ci==7 else "Arial"),size=(9 if ci==7 else 9.5),
                        color=("C0504D" if (ci==11 and st=="Unused calc") else "000000"))
            c.alignment=Alignment(vertical="top",wrap_text=True); c.border=border
            if fill!="FFFFFF": c.fill=PatternFill("solid",fgColor=fill)
        rn+=1
    d.freeze_panes=f"A{hdr+1}"; d.auto_filter.ref=f"A{hdr}:{get_column_letter(len(cols))}{rn-1}"

    # Dashboard actions
    _actions_sheet(book, wb["actions"])

    # Lint / health-check findings (embedded from lint_workbook)
    add_findings_sheet(book, findings, "Lint")

    if not a.no_graph:
        try:
            from dependency_graph import render_graph, add_graph_sheet
            png = os.path.splitext(out)[0] + "_graph.png"
            n_nodes, n_edges = render_graph(fields, png)
            if n_edges:
                add_graph_sheet(book, png)
                print(f"Dependency graph: {n_nodes} nodes, {n_edges} edges (sheet added; image: {png})")
            else:
                print("Dependency graph skipped: no calculated-field dependencies to draw.")
        except ImportError:
            print("Dependency graph skipped: matplotlib not installed (pip install matplotlib).")

    book.save(out)
    print(f"Data dictionary written: {out}  ({len(fields)} fields)")

if __name__=="__main__":
    main()
