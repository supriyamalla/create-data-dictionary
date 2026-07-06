"""
Dependency-graph renderer for the data dictionary.

Draws a left-to-right graph of how fields feed calculated fields, and adds it as a
"Dependency Graph" sheet to the workbook. Layout (longest-path layering + bidirectional
barycentre ordering) is computed here, so it needs only matplotlib - no Graphviz/networkx.

Imported by build_dictionary.py when --graph is passed. Can also run standalone on a
parse_workbook.py fields.json for iterating on the visual:

  python dependency_graph.py fields.json --png graph.png
"""
import textwrap
from collections import defaultdict

# palette: (fill, edge) by node kind / status
COL_PARAM  = ("#FFF3C4", "#C99700")   # parameters
COL_SOURCE = ("#D6E4FF", "#3A66B0")   # raw datasource columns
COL_CALC   = ("#CDEBE3", "#2E8B7F")   # calculated, in use
COL_UNUSED = ("#FBD5D5", "#C0392B")   # calculated, unused (dead logic)

# above this many connected nodes the graph gets dense; we still draw it but warn.
DENSE_NODE_WARNING = 60


def build_graph(fields):
    """fields -> (by_caption, node_captions, edges[(parent, child)])."""
    by_cap = {f["caption"]: f for f in fields}
    edges, participating = [], set()
    for f in fields:
        for dep in (f.get("dependencies") or []):
            if dep in by_cap:
                edges.append((dep, f["caption"]))
                participating.update((dep, f["caption"]))
    nodes = [c for c in by_cap if c in participating]
    return by_cap, nodes, edges


def _layer(nodes, edges):
    parents, children = defaultdict(list), defaultdict(list)
    for a, b in edges:
        parents[b].append(a)
        children[a].append(b)

    level = {}
    def lvl(n, seen=()):                       # longest path from a root
        if n in level:
            return level[n]
        if n in seen or not parents[n]:
            level[n] = 0
            return 0
        level[n] = 1 + max(lvl(p, seen + (n,)) for p in parents[n])
        return level[n]
    for n in nodes:
        lvl(n)

    cols = defaultdict(list)
    for n in nodes:
        cols[level[n]].append(n)

    order = {n: i for c in cols for i, n in enumerate(sorted(cols[c]))}
    def sweep(column_ids, neighbours):
        for c in column_ids:
            def bary(n):
                nb = neighbours[n]
                return sum(order[m] for m in nb) / len(nb) if nb else order[n]
            cols[c].sort(key=bary)
            for i, n in enumerate(cols[c]):
                order[n] = i
    asc = sorted(cols)
    for _ in range(8):                         # bidirectional crossing reduction
        sweep(asc[1:], parents)
        sweep(asc[-2::-1], children)
    return level, cols


def _style(f):
    if f["kind"] == "parameter":
        return COL_PARAM
    if f["kind"] == "field":
        return COL_SOURCE
    return COL_UNUSED if f.get("unused") else COL_CALC


def render_graph(fields, png_path):
    """Render the dependency graph to png_path. Returns (n_nodes, n_edges); (0,0) if nothing to draw."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Patch

    by_cap, nodes, edges = build_graph(fields)
    if not edges:
        return 0, 0
    if len(nodes) > DENSE_NODE_WARNING:
        print(f"  note: {len(nodes)} connected fields - the graph will be dense; "
              f"the Summary sheet remains the authoritative field list.")

    level, cols = _layer(nodes, edges)
    max_lvl = max(cols)
    max_rows = max(len(v) for v in cols.values())

    XW, YH = 4.6, 1.35          # column pitch, row pitch
    BW, BH = 3.7, 0.92          # box width / height
    pos = {}
    for c, ns in cols.items():
        y_off = (max_rows - len(ns)) / 2.0
        for i, n in enumerate(ns):
            pos[n] = (c * XW, (max_rows - (i + y_off)) * YH)

    fig, ax = plt.subplots(figsize=(max(12, (max_lvl + 1) * XW * 0.9),
                                    max(7, max_rows * YH * 0.72)))
    for a, b in edges:
        x1, y1 = pos[a]; x2, y2 = pos[b]
        ax.add_patch(FancyArrowPatch(
            (x1 + BW / 2, y1), (x2 - BW / 2, y2),
            connectionstyle="arc3,rad=0.06", arrowstyle="-|>",
            mutation_scale=11, lw=0.9, color="#9AA4B0", alpha=0.55, zorder=1))
    for n in nodes:
        x, y = pos[n]
        fill, edge = _style(by_cap[n])
        ax.add_patch(FancyBboxPatch(
            (x - BW / 2, y - BH / 2), BW, BH,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            fc=fill, ec=edge, lw=1.4, zorder=2))
        ax.text(x, y, "\n".join(textwrap.wrap(n, 24)) or n,
                ha="center", va="center", fontsize=8.2, color="#1A1A1A", zorder=3)

    handles = [Patch(fc=COL_SOURCE[0], ec=COL_SOURCE[1], label="Source column"),
               Patch(fc=COL_PARAM[0],  ec=COL_PARAM[1],  label="Parameter"),
               Patch(fc=COL_CALC[0],   ec=COL_CALC[1],   label="Calculated (in use)"),
               Patch(fc=COL_UNUSED[0], ec=COL_UNUSED[1], label="Calculated (unused)")]
    ax.legend(handles=handles, loc="lower center", ncol=4, frameon=False,
              fontsize=9, bbox_to_anchor=(0.5, -0.04))
    ax.set_title("Calculated-field dependency graph  (data flows left → right)",
                 fontsize=13, fontweight="bold", pad=14)
    ax.autoscale(); ax.margins(0.04); ax.axis("off")
    fig.tight_layout()
    fig.savefig(png_path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return len(nodes), len(edges)


def add_graph_sheet(book, png_path):
    """Add a 'Dependency Graph' sheet holding png_path to an open openpyxl Workbook."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font
    if "Dependency Graph" in book.sheetnames:
        del book["Dependency Graph"]
    ws = book.create_sheet("Dependency Graph")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Dependency Graph"; ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = ("Each arrow points from a field to the calculation that uses it. "
                "Red boxes are calculations nothing depends on (candidate dead logic). "
                "Fields with no connections are omitted here - see the Summary sheet for the full list.")
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="555555")
    ws.add_image(XLImage(png_path), "A4")


if __name__ == "__main__":
    import argparse, json
    ap = argparse.ArgumentParser(description="Render dependency graph from a fields.json")
    ap.add_argument("fields_json")
    ap.add_argument("--png", default="dependency_graph.png")
    a = ap.parse_args()
    fields = json.load(open(a.fields_json, encoding="utf-8"))["fields"]
    n, e = render_graph(fields, a.png)
    print(f"Graph: {n} nodes, {e} edges -> {a.png}" if e else "No calculated dependencies to graph.")
