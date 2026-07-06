"""
Lint / health-check a Tableau workbook.

Reads the same parsed model as the data dictionary and reports maintainability
and governance issues: dead calculated fields, calcs missing a description,
high-complexity (LOD / table-calc / deeply-nested) calcs, auto-generated field
names that were never renamed, and captions that collide across datasources.

Usage:
  python lint_workbook.py WORKBOOK.twbx [-o REPORT.md]

Findings are printed to the console and written to a Markdown report (browsable
in Obsidian). This is read-only — it never modifies the workbook.
"""
import argparse, os, re
from tableau_xml import parse_workbook

# table-calculation functions — matched only as real function calls (name followed
# by "(") so the words don't false-positive inside strings or field names.
TABLE_CALCS = (r"\b(INDEX|SIZE|FIRST|LAST|LOOKUP|PREVIOUS_VALUE|TOTAL|"
               r"RANK|RANK_DENSE|RANK_MODIFIED|RANK_PERCENTILE|RANK_UNIQUE|"
               r"RUNNING_(SUM|AVG|MIN|MAX|COUNT)|WINDOW_[A-Z]+)\s*\(")
LOD = r"\{\s*(FIXED|INCLUDE|EXCLUDE)\b"


def complexity_flags(formula):
    """Return the list of complexity/risk tags that apply to a formula."""
    flags = []
    body = re.sub(r'"[^"]*"|\'[^\']*\'', "", formula)  # ignore string literals
    if re.search(LOD, body, re.I):
        flags.append("LOD")
    if re.search(TABLE_CALCS, body, re.I):
        flags.append("table calc")
    branches = len(re.findall(r"\b(if|elseif|when)\b", body, re.I))
    if branches >= 3:
        flags.append(f"{branches}-way branch")
    if len(formula) > 200:
        flags.append(f"{len(formula)} chars")
    return flags


def auto_named(caption):
    """A calc Tableau auto-named and the author never renamed."""
    return bool(re.match(r"(Calculation_?\d|Calculation\d*$)", caption))


# check registry: (report/Excel label, findings key, why-it-matters, Excel row colour)
CHECKS = [
    ("Dead calculated field", "dead",
     "On no sheet and feeding no other field — safe-to-remove candidate.", "FCE4D6"),
    ("Missing description", "no_desc",
     "No in-workbook description; meaning lives only in the formula.", "DDEBF7"),
    ("High complexity", "complex",
     "LOD / table-calc / deeply-nested / long — review for correctness & performance.", "FFF2CC"),
    ("Auto-generated name", "auto_name",
     "Left as Tableau's default Calculation_… name — rename for clarity.", "FCE4D6"),
    ("Caption reused across datasources", "dupes",
     "Same name in more than one datasource — ambiguous in conversation.", "E2EFDA"),
]


def suggest_fix(key, flags):
    """A concrete remediation hint. For complexity it is tailored to the flags present."""
    if key == "dead":
        return "If genuinely unused, delete it; otherwise place it on a sheet or reference it in another calc."
    if key == "no_desc":
        return ("Add a description in Tableau (right-click the field → Default Properties → Comment), "
                "or supply one via descriptions.json so the dictionary carries it.")
    if key == "auto_name":
        return "Rename to a business-meaningful name so it reads clearly everywhere it's used."
    if key == "dupes":
        return "Rename one of the fields, or always qualify it by datasource to avoid ambiguity."
    if key == "complex":
        tips = []
        for fl in flags:
            if fl == "LOD":
                tips.append("FIXED LODs ignore dimension/quick filters — confirm that's intended (use a "
                            "context filter, or switch to INCLUDE/EXCLUDE, if not); for a fixed grain, "
                            "consider precomputing it in the data source or extract.")
            elif fl == "table calc":
                tips.append("Table calcs depend on Compute-Using / partitioning and can silently break when "
                            "the view changes — document the addressing; if the result is order-independent, "
                            "an LOD or a plain aggregation is more robust.")
            elif fl.endswith("branch"):
                tips.append("Simplify the branching: replace IF/ELSEIF chains with CASE where possible, and "
                            "push value-to-value mappings into a lookup/join table, a group, or a parameter "
                            "rather than a formula; extract any repeated sub-expressions into their own fields.")
            elif fl.endswith("chars"):
                tips.append("Split into smaller intermediate calculated fields for readability and reuse.")
        return " ".join(dict.fromkeys(tips))  # de-dupe, preserve order
    return ""


def finding_rows(key, items):
    """Normalise a check's findings to dicts: field, datasource, detail, formula, fix."""
    rows = []
    for it in items:
        if key == "complex":
            f, flags = it
            rows.append({"field": f["caption"], "datasource": f["datasource"],
                         "detail": ", ".join(flags), "formula": f["formula"],
                         "fix": suggest_fix(key, flags)})
        elif key == "dupes":
            cap, dss = it
            rows.append({"field": cap, "datasource": ", ".join(dss),
                         "detail": f"in {len(dss)} datasources", "formula": "",
                         "fix": suggest_fix(key, [])})
        else:
            rows.append({"field": it["caption"], "datasource": it["datasource"],
                         "detail": "", "formula": it["formula"], "fix": suggest_fix(key, [])})
    return rows


def lint(wb):
    calcs = [f for f in wb["fields"] if f["kind"] == "calculated"]
    findings = {}

    findings["dead"] = [f for f in calcs if f["unused"]]
    findings["no_desc"] = [f for f in calcs if not f["description"]]
    findings["complex"] = [(f, complexity_flags(f["formula"])) for f in calcs
                           if complexity_flags(f["formula"])]
    findings["auto_name"] = [f for f in calcs if auto_named(f["caption"])]

    seen = {}
    for f in wb["fields"]:
        seen.setdefault(f["caption"], set()).add(f["datasource"])
    findings["dupes"] = sorted(((c, sorted(ds)) for c, ds in seen.items() if len(ds) > 1),
                               key=lambda x: x[0])

    return calcs, findings


def build_report(wb, calcs, findings):
    L = [f"# Lint report — {wb['source']}\n",
         f"{len(wb['datasources'])} datasources · {len(wb['worksheets'])} worksheets · "
         f"{len(wb['fields'])} fields ({len(calcs)} calculated)\n",
         "## Summary\n", "| Check | Count |", "|---|---|"]
    for label, key, _why, _fill in CHECKS:
        L.append(f"| {label} | {len(findings[key])} |")
    L.append("")
    for label, key, why, _fill in CHECKS:
        rows = finding_rows(key, findings[key])
        L.append(f"## {label}  ({len(rows)})\n")
        L.append(f"_{why}_\n")
        if not rows:
            L.append("- none\n")
            continue
        for row in rows:
            suffix = f" — {row['detail']}" if row["detail"] else ""
            L.append(f"- **{row['field']}** ({row['datasource']}){suffix}")
            if key == "complex" and row["formula"]:
                L.append(f"  - Formula: `{row['formula']}`")
            if row["fix"]:
                L.append(f"  - Suggested fix: {row['fix']}")
        L.append("")
    return "\n".join(L)


def write_excel(wb, calcs, findings, out):
    """Write the same findings as a two-sheet .xlsx (Summary + filterable Findings)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side
    NAVY, HEAD = "1F3864", "2F5496"
    arial = lambda **k: Font(name="Arial", **k)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    book = Workbook()

    # Summary sheet
    s = book.active; s.title = "Lint Summary"
    s["A1"] = "Tableau Workbook — Lint / Health Check"
    s["A1"].font = arial(bold=True, size=15, color="FFFFFF")
    s["A1"].fill = PatternFill("solid", fgColor=NAVY); s.merge_cells("A1:B1")
    s.row_dimensions[1].height = 24
    s["A2"] = wb["source"]; s["A2"].font = arial(italic=True, size=10, color="555555")
    s["A3"] = (f"{len(wb['datasources'])} datasources · {len(wb['worksheets'])} worksheets · "
               f"{len(wb['fields'])} fields ({len(calcs)} calculated)")
    s["A3"].font = arial(size=10, color="555555")
    s["A5"] = "Check"; s["B5"] = "Count"
    for cell in ("A5", "B5"):
        s[cell].font = arial(bold=True, color="FFFFFF", size=10)
        s[cell].fill = PatternFill("solid", fgColor=HEAD)
    r = 6
    for label, key, _why, fill in CHECKS:
        s[f"A{r}"] = label; s[f"B{r}"] = len(findings[key])
        s[f"A{r}"].font = arial(size=10); s[f"B{r}"].font = arial(size=10, bold=True)
        s[f"A{r}"].fill = PatternFill("solid", fgColor=fill)
        s[f"A{r}"].border = border; s[f"B{r}"].border = border
        r += 1
    s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 10

    add_findings_sheet(book, findings, "Findings")
    book.save(out)


def add_findings_sheet(book, findings, sheet_name="Findings"):
    """Add a filterable, colour-banded findings sheet to an existing openpyxl book.

    Shared by the standalone lint report and the data dictionary (which embeds it
    as a 'Lint' sheet)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HEAD = "2F5496"
    arial = lambda **k: Font(name="Arial", **k)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    d = book.create_sheet(sheet_name)
    cols = ["Check", "Field", "Datasource", "Detail", "Formula", "Suggested fix"]
    for i, w in enumerate([22, 26, 24, 18, 50, 62], 1):
        d.column_dimensions[get_column_letter(i)].width = w
        c = d.cell(row=1, column=i, value=cols[i - 1])
        c.font = arial(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    d.row_dimensions[1].height = 22
    rn = 2
    for label, key, _why, fill in CHECKS:
        for row in finding_rows(key, findings[key]):
            d.append([label, row["field"], row["datasource"], row["detail"],
                      row["formula"], row["fix"]])
            for ci in range(1, 7):
                c = d.cell(row=rn, column=ci)
                mono = (ci == 5)
                c.font = Font(name="Consolas" if mono else "Arial", size=9 if mono else 9.5)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = border; c.fill = PatternFill("solid", fgColor=fill)
            rn += 1
    if rn == 2:
        d.append(["— no findings —"])
    d.freeze_panes = "A2"; d.auto_filter.ref = f"A1:F{max(rn - 1, 1)}"
    return d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("-o", "--output", help="report path; .xlsx writes Excel, otherwise Markdown")
    a = ap.parse_args()

    wb = parse_workbook(a.workbook)
    calcs, findings = lint(wb)
    out = a.output or os.path.splitext(os.path.basename(a.workbook))[0] + "_lint.md"

    total = sum(len(findings[k]) for k in findings)
    print(f"{wb['source']}: {len(calcs)} calculated fields, {total} findings")
    for label, key in [("dead code", "dead"), ("missing description", "no_desc"),
                       ("high complexity", "complex"), ("auto-named", "auto_name"),
                       ("duplicate captions", "dupes")]:
        print(f"  {label:22} {len(findings[key])}")

    if out.lower().endswith(".xlsx"):
        write_excel(wb, calcs, findings, out)
    else:
        with open(out, "w", encoding="utf-8") as f:
            f.write(build_report(wb, calcs, findings))
    print(f"Lint report written: {out}")


if __name__ == "__main__":
    main()
