"""
Style / design QC for a Tableau workbook.

Reads the same .twb XML as the dictionary and checks the things a human
reviewer checks by eye before a dashboard ships: consistent fonts, a consistent
colour scheme, and consistent number formatting.

Two modes, combined:
  * majority-wins  - infers the convention from the workbook itself ("14 sheets
    use Tableau Book 10pt, 1 uses Tahoma 9pt"). Needs no configuration, so it
    runs on any workbook, including a client's.
  * style guide    - an optional JSON of approved fonts / colours / number
    formats, so it can also flag a workbook that is consistently off-brand.

Usage:
  python style_check.py WORKBOOK.twbx [-g style_guide.json]

Read-only - it never modifies the workbook.
"""
import argparse, json, os, re, sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from tableau_xml import load_twb_text

# check registry: (label, findings key, why-it-matters, Excel row colour)
STYLE_CHECKS = [
    ("Font family", "font_family",
     "A sheet uses a different typeface from the rest of the workbook.", "DDEBF7"),
    ("Font size", "font_size",
     "Same element styled at a different point size elsewhere.", "DDEBF7"),
    ("Number format", "number_format",
     "Fields of the same kind shown at different precision or scale.", "FFF2CC"),
    ("Off-palette colour", "color",
     "A colour outside the approved scheme, or used in only one or two places.", "FCE4D6"),
]

# element roles worth checking - the chrome a reviewer actually looks at.
FONT_ROLES = {"worksheet", "title", "header", "axis", "legend", "legend-title-text",
              "quick-filter", "cell", "label", "dash-title", "dash-subtitle"}
# <run> overrides live under these holders; tooltips are deliberately out of scope.
RUN_ROLES = {"title": "title", "caption": "title", "customized-label": "label",
             "formatted-text": "title"}
# a majority is only a standard if enough places voted and it actually won
MIN_PLACES = 3
COLOR_ATTRS = {"color", "background-color", "border-color", "stroke-color", "mark-color"}
HEXCOLOR = re.compile(r"^#(?:[0-9a-f]{6}|[0-9a-f]{8})$", re.I)
# greys / black / white / transparent are Tableau defaults, not scheme choices
NEUTRAL = re.compile(r"^#(?:[0-9a-f]{2})?(?:([0-9a-f]{2})\1\1)$", re.I)

_DATE_TOK = re.compile(r"(yyyy|yy|mmmm|mmm|dd|hh|ss)", re.I)
_CORE = re.compile(r"[#0][#0,]*(?:\.[#0]+)?,*")


def classify_format(code):
    """Tableau format code -> (family, decimals, scale_commas). None if unparseable.

    Examples: 'p0.0%' -> ('percent', 1, 0);  'c"$"#,##0,K' -> ('currency', 0, 1);
              'n#,##0.0"%"' -> ('percent', 1, 0)  (a percent faked with a number format)
    """
    if not code:
        return None
    body = code.split(";")[0]                       # positive section only
    lead = ""
    if body[:1] in "pcn":
        lead, body = body[0], body[1:]
    else:
        m = re.match(r"^[Cc](\d+)", body)           # locale-default, e.g. C1033%
        if m:
            lead, body = "C", body[m.end():]
    lits = re.findall(r'"([^"]*)"', body)
    stripped = re.sub(r'"[^"]*"', "", body)
    blob = stripped + "".join(lits)

    if _DATE_TOK.search(stripped) and "#" not in stripped:
        return ("date", None, None)
    if "%" in blob:
        family = "percent"
    elif lead in ("c", "C") or re.search(r"[$\u20ac\u00a3\u00a5]", blob):
        family = "currency"
    elif lead == "n" or "#" in stripped or "0" in stripped:
        family = "number"
    else:
        return None

    m = _CORE.search(stripped)
    if not m:
        return (family, None, None)
    core = m.group(0)
    scale = len(core) - len(core.rstrip(","))
    dec = core.rstrip(",")
    decimals = len(dec.split(".", 1)[1]) if "." in dec else 0
    return (family, decimals, scale)


def describe_format(cls):
    """Human wording for a classified format tuple."""
    family, decimals, scale = cls
    if family == "date":
        return "date"
    bits = [family]
    if decimals is not None:
        bits.append(f"{decimals} decimal" + ("" if decimals == 1 else "s"))
    if scale:
        bits.append("scaled by " + {1: "thousands", 2: "millions"}.get(scale, f"1000^{scale}"))
    return ", ".join(bits)


def parse_styles(path):
    """Collect every font, colour and number-format decision, with its scope."""
    text, _ = load_twb_text(path)
    root = ET.fromstring(text)
    pmap = {c: p for p in root.iter() for c in p}

    def scope_of(node):
        """(container kind, container name, style-rule element role) for a node."""
        role = ""
        n = node
        while n in pmap:
            n = pmap[n]
            if not role and n.tag == "style-rule":
                role = n.get("element") or ""
            if n.tag in ("worksheet", "dashboard") and n.get("name"):
                return n.tag, n.get("name"), role
        return "workbook", "(workbook)", role

    fonts, colors = [], []
    for fmt in root.iter("format"):
        attr, val = fmt.get("attr"), fmt.get("value")
        if not attr or val is None:
            continue
        kind, name, role = scope_of(fmt)
        if attr in ("font-family", "font-size"):
            if role and role not in FONT_ROLES:
                continue
            fonts.append({"scope": name, "scope_kind": kind, "role": role or "worksheet",
                          "attr": attr, "value": val})
        elif attr in COLOR_ATTRS and HEXCOLOR.match(val.strip()):
            colors.append({"scope": name, "scope_kind": kind, "role": role or "worksheet",
                           "attr": attr, "value": val.lower()})

    # title / caption runs carry their own font overrides
    for run in root.iter("run"):
        kind, name, role = scope_of(run)
        holder = pmap.get(pmap.get(run))          # formatted-text -> title/caption
        role = role or RUN_ROLES.get(holder.tag if holder is not None else "title", "")
        if role not in FONT_ROLES:
            continue
        if run.get("fontname"):
            fonts.append({"scope": name, "scope_kind": kind, "role": role,
                          "attr": "font-family", "value": run.get("fontname")})
        if run.get("fontsize"):
            fonts.append({"scope": name, "scope_kind": kind, "role": role,
                          "attr": "font-size", "value": run.get("fontsize")})

    palettes = []
    for p in root.iter("color-palette"):
        palettes.append({"name": p.get("name") or "(unnamed)", "type": p.get("type"),
                         "colors": [(c.text or "").strip().lower() for c in p if c.text]})

    formats = []
    seen = set()
    for col in root.findall("./datasources/datasource/column"):
        code = col.get("default-format")
        if not code:
            continue
        cls = classify_format(code)
        if not cls or cls[0] == "date":
            continue
        cap = col.get("caption") or (col.get("name") or "").strip("[]")
        if (cap, code) in seen:
            continue
        seen.add((cap, code))
        formats.append({"caption": cap, "code": code, "cls": cls})

    return {"fonts": fonts, "colors": colors, "palettes": palettes, "formats": formats,
            "worksheets": [w.get("name") for w in root.findall(".//worksheets/worksheet")]}


def _f(check, where, element, found, standard, fix):
    return {"check": check, "where": where, "element": element,
            "found": found, "standard": standard, "fix": fix}


def check_styles(styles, guide=None):
    """Return {check_key: [finding, ...]} from majority-wins plus an optional style guide."""
    out = {k: [] for _, k, _, _ in STYLE_CHECKS}
    guide = guide or {}

    # ---- fonts: per element role, the value most places use is the standard ----
    uniq = {(o["scope"], o["role"], o["attr"], o["value"]) for o in styles["fonts"]}
    groups = defaultdict(list)
    for scope, role, attr, value in uniq:
        groups[(role, attr)].append((scope, value))
    approved_fonts = guide.get("fonts") or {}
    for (role, attr), obs in sorted(groups.items()):
        counts = Counter(v for _s, v in obs)
        major, n_major = counts.most_common(1)[0]
        if len(obs) < MIN_PLACES or n_major * 2 <= len(obs):
            continue          # no defensible majority - stay quiet rather than guess
        key = "font_family" if attr == "font-family" else "font_size"
        ok = approved_fonts.get("family" if attr == "font-family" else "size")
        unit = "pt" if attr == "font-size" else ""
        nice = attr.replace("-", " ")
        for scope, value in sorted(obs):
            off_guide = ok is not None and value not in ok
            if value == major and not off_guide:
                continue
            if off_guide:
                std = "approved: " + ", ".join(str(x) + unit for x in ok)
                fix = "Set " + role + " " + nice + " to an approved value."
            else:
                std = str(major) + unit + " (" + str(n_major) + " of " + str(len(obs)) + " places)"
                fix = "Change " + role + " " + nice + " to " + str(major) + unit + "."
            out[key].append(_f("Font family" if key == "font_family" else "Font size",
                               scope, role, str(value) + unit, std, fix))

    # ---- number formats: within a family, the most common precision is the standard ----
    byfam = defaultdict(list)
    for f in styles["formats"]:
        byfam[f["cls"][0]].append(f)
    want = guide.get("number_formats") or {}
    for family, items in sorted(byfam.items()):
        # only fields with an explicit precision vote on the standard; a locale
        # default ("use whatever the locale does") is not a precision decision.
        items = [f for f in items if f["cls"][1] is not None]
        if not items:
            continue
        counts = Counter(f["cls"] for f in items)
        major, n_major = counts.most_common(1)[0]
        g = want.get(family)
        if g is None and len(counts) > 1 and n_major * 2 <= len(items):
            # genuinely no house standard - say that once, rather than pick a
            # winner that most of the workbook disagrees with.
            variants = "; ".join(describe_format(c) + " x" + str(n)
                                 for c, n in counts.most_common())
            out["number_format"].append(
                _f("Number format", str(len(items)) + " " + family + " fields", family,
                   str(len(counts)) + " different formats in use",
                   "no consistent standard - " + variants,
                   "Agree one " + family + " format and apply it to all "
                   + str(len(items)) + " fields."))
            continue
        for f in sorted(items, key=lambda x: x["caption"]):
            off_guide = g is not None and f["cls"][1] != g.get("decimals")
            if f["cls"] == major and not off_guide:
                continue
            if off_guide:
                dp = g.get("decimals")
                std = ("approved: " + family + ", " + str(dp) + " decimal"
                       + ("" if dp == 1 else "s"))
                fix = "Reformat to " + str(g.get("decimals")) + " decimal places."
            else:
                std = (describe_format(major) + " (" + str(n_major) + " of "
                       + str(len(items)) + " fields)")
                fix = "Reformat to match: " + describe_format(major) + "."
            out["number_format"].append(
                _f("Number format", f["caption"], family,
                   describe_format(f["cls"]) + "  (" + f["code"] + ")", std, fix))

    # ---- colours ----
    spots = {(c["scope"], c["role"], c["attr"], c["value"]) for c in styles["colors"]}
    live = [s for s in spots if not NEUTRAL.match(s[3]) and s[3] != "#00000000"]
    approved = {c.lower() for c in (guide.get("colors") or [])}
    if approved:
        for scope, role, attr, value in sorted(live):
            if value not in approved:
                out["color"].append(
                    _f("Off-palette colour", scope, role + " · " + attr, value,
                       "approved scheme (" + str(len(approved)) + " colours)",
                       "Replace with the nearest approved brand colour."))
    elif len(live) >= 5:
        usage = Counter(v for _s, _r, _a, v in live)
        for scope, role, attr, value in sorted(live):
            if usage[value] <= 2:
                plural = "" if usage[value] == 1 else "s"
                out["color"].append(
                    _f("Off-palette colour", scope, role + " · " + attr, value,
                       "used in only " + str(usage[value]) + " place" + plural,
                       "Confirm this is intentional, or fold into the scheme."))
    return out


def summary_counts(findings):
    """[(label, count), ...] in registry order - for the dictionary's Summary sheet."""
    return [(label, len(findings[key])) for label, key, _why, _fill in STYLE_CHECKS]


def add_style_sheet(book, findings, sheet_name="Style"):
    """Add a filterable, colour-banded design-QC sheet to an existing openpyxl book."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from lint_workbook import legend_block
    HEAD = "2F5496"
    arial = lambda **k: Font(name="Arial", **k)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sh = book.create_sheet(sheet_name)
    cols = ["Check", "Where", "Element", "Found", "Workbook standard", "Suggested fix"]
    for i, w in enumerate([18, 30, 24, 34, 34, 46], 1):
        sh.column_dimensions[get_column_letter(i)].width = w
    bycolor = {}
    for label, key, _why, fill in STYLE_CHECKS:
        bycolor.setdefault(fill, []).append(label)
    hdr = legend_block(sh, [(fill, " / ".join(labs)) for fill, labs in bycolor.items()])
    for i in range(1, len(cols) + 1):
        c = sh.cell(row=hdr, column=i, value=cols[i - 1])
        c.font = arial(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    sh.row_dimensions[hdr].height = 22
    rn = hdr + 1
    for label, key, _why, fill in STYLE_CHECKS:
        for r in findings[key]:
            vals = [label, r["where"], r["element"], r["found"], r["standard"], r["fix"]]
            for ci in range(1, len(cols) + 1):
                c = sh.cell(row=rn, column=ci, value=vals[ci - 1])
                c.font = arial(size=9.5)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = border
                c.fill = PatternFill("solid", fgColor=fill)
            rn += 1
    if rn == hdr + 1:
        sh.cell(row=rn, column=1, value="— no style findings —")
        rn += 1
    sh.freeze_panes = "A" + str(hdr + 1)
    sh.auto_filter.ref = ("A" + str(hdr) + ":" + get_column_letter(len(cols))
                          + str(max(rn - 1, hdr)))
    return sh


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description="Design-QC a Tableau workbook (read-only).")
    ap.add_argument("workbook")
    ap.add_argument("-g", "--style-guide", help="optional JSON of approved fonts/colours/formats")
    a = ap.parse_args()
    guide = json.load(open(a.style_guide, encoding="utf-8")) if a.style_guide else None
    styles = parse_styles(a.workbook)
    findings = check_styles(styles, guide)

    mode = "  [style guide applied]" if guide else "  [majority-wins]"
    print("\nStyle check — " + os.path.basename(a.workbook))
    print("  " + str(len(styles["worksheets"])) + " sheets | "
          + str(len(styles["formats"])) + " formatted fields | "
          + str(len(styles["palettes"])) + " custom palette(s)" + mode)
    total = 0
    for label, key, why, _fill in STYLE_CHECKS:
        rows = findings[key]
        total += len(rows)
        print("\n" + label + ": " + str(len(rows)))
        if rows:
            print("  " + why)
        for r in rows[:8]:
            print("  - " + r["where"] + " | " + r["element"] + ": "
                  + r["found"] + "  ->  " + r["standard"])
        if len(rows) > 8:
            print("  ... and " + str(len(rows) - 8) + " more")
    print("\nTotal: " + str(total) + " finding(s)\n")


if __name__ == "__main__":
    main()
